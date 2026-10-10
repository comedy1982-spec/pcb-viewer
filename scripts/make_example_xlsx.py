"""Writes examples/SENSE-32_HSIS.xlsx from the built-in sample HSIS (needs openpyxl).
Run through `npm run examples`, which passes the sample as JSON."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1], encoding='utf-8'))
out = sys.argv[2] if len(sys.argv) > 2 else 'examples/SENSE-32_HSIS.xlsx'
wb = Workbook()
info = wb.active
info.title = '문서 정보'
m = data['meta']
for k, v in [('문서 ID', m['docId']), ('제목', m['title']), ('버전', m['version']), ('MCU 참조', m['mcuRef']),
             ('작성일', m['date']), ('대상 보드', 'SENSE-32 DEV v1.2'), ('ASPICE 산출물', 'HW/SW 인터페이스 사양 (HWE.1 / SWE.1)')]:
    info.append([k, v])
for r in info.iter_rows(min_col=1, max_col=1):
    r[0].font = Font(bold=True)
info.column_dimensions['A'].width = 16
info.column_dimensions['B'].width = 44

ws = wb.create_sheet('HSIS')
ws.append(data['header'])
for row in data['rows']:
    ws.append(row)
head_fill = PatternFill('solid', fgColor='1E3A4C')
thin = Side(style='thin', color='C9D1D6')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = head_fill
    c.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=1):
    for c in row:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if c.row > 1:
            c.alignment = Alignment(vertical='top', wrap_text=True)
widths = {'ID': 11, '신호명': 14, '분류': 11, 'MCU 핀': 14, '포트': 8, '핀 기능': 13, '주변장치': 9, '방향': 12, '초기값': 7, '활성 레벨': 8,
          '풀업/풀다운': 15, '회로도 넷': 12, 'HW 연결': 18, '전기적 특성': 20, '변환식': 26, '타이밍': 14, 'SW 모듈': 9, 'SW 설정 ID': 28,
          'API': 18, '요구사항 ID': 13, '테스트 ID': 12, 'ASIL': 6, '상태': 8, '비고': 24}
for i, h in enumerate(data['header'], 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions
wb.active = 1
wb.save(out)
print(out)
